"""Colokan #2: scheduled export folder.

RME lama dump CSV/XLSX ke sebuah folder (atau SFTP yang sudah ter-mount) tiap
jam/hari; MANTRA baca dari situ. Paling universal — hampir semua RME bisa
export file walau tak punya API/DB terbuka. Read-only: hanya membaca, file
sumber tak diubah (arsip/rename diserahkan ke pipeline setelah audit).

Konvensi: satu entitas = satu file, dinamai `<entity>.csv` / `<entity>.xlsx`
di dalam `base_dir`. Contoh: `patient.csv`, `encounter.csv`.
"""

from __future__ import annotations

import csv
from pathlib import Path

from sentra_mantra_integrations.rme_bridge.sources.base import Source


class FolderSource(Source):
	kind = "folder"

	def __init__(self, base_dir: str) -> None:
		self.base_dir = Path(base_dir)

	def _resolve(self, entity: str) -> Path:
		"""Cari file export untuk `entity` di folder; utamakan CSV lalu XLSX."""
		for ext in (".csv", ".xlsx"):
			candidate = self.base_dir / f"{entity}{ext}"
			if candidate.exists():
				return candidate
		raise FileNotFoundError(
			f"Tidak ada file export '{entity}.csv/.xlsx' di {self.base_dir}"
		)

	def fetch(self, entity: str) -> list[dict]:
		path = self._resolve(entity)
		if path.suffix == ".csv":
			return self._read_csv(path)
		return self._read_xlsx(path)

	@staticmethod
	def _read_csv(path: Path) -> list[dict]:
		with path.open(newline="", encoding="utf-8-sig") as fh:
			return [dict(row) for row in csv.DictReader(fh)]

	@staticmethod
	def _read_xlsx(path: Path) -> list[dict]:
		# openpyxl dibawa Frappe; import lokal supaya modul tetap ringan saat CSV.
		from openpyxl import load_workbook

		wb = load_workbook(path, read_only=True, data_only=True)
		ws = wb.active
		rows = ws.iter_rows(values_only=True)
		try:
			header = [str(c) if c is not None else "" for c in next(rows)]
		except StopIteration:
			return []
		out: list[dict] = []
		for values in rows:
			out.append({header[i]: values[i] for i in range(len(header))})
		return out

	def describe(self) -> dict:
		return {"kind": self.kind, "base_dir": str(self.base_dir)}
