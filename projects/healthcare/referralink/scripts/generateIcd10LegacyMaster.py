from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) != 3:
        print(
            "Usage: python scripts/generateIcd10LegacyMaster.py <input.xlsx> <output.json>",
            file=sys.stderr,
        )
        return 1

    source_path = Path(sys.argv[1]).expanduser().resolve()
    output_path = Path(sys.argv[2]).expanduser().resolve()

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if "ICD10" not in workbook.sheetnames:
        print("Workbook is missing the expected 'ICD10' sheet.", file=sys.stderr)
        return 1

    sheet = workbook["ICD10"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        print("Workbook sheet is empty.", file=sys.stderr)
        return 1

    expected_header = ("CODE", "DISPLAY", "VERSION")
    if tuple(header) != expected_header:
        print(f"Unexpected header: {header!r}", file=sys.stderr)
        return 1

    entries = []
    versions = set()

    for row in rows:
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue

        code, display, version = (str(value).strip() for value in row)
        versions.add(version)
        entries.append(
            {
                "code": code,
                "officialLabel": display,
                "version": version,
            }
        )

    document = {
        "source": "ICD-10 e-klaim workbook",
        "sourcePath": str(source_path),
        "sheetName": "ICD10",
        "rowCount": len(entries),
        "versions": sorted(versions),
        "entries": entries,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(document, ensure_ascii=True, indent=2) + "\n",
        encoding="utf-8",
    )

    print(
        f"Wrote {len(entries)} ICD master entries from {source_path.name} to {output_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
